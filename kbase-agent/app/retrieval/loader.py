"""文档解析：把 data/docs 下支持的格式统一抽取为纯文本/Markdown。

支持：.md / .txt（UTF-8 直读）、.docx / .xlsx / .pdf（文本层抽取）。
文件名以 `_` 开头视为说明文件，不参与建索引。
source 用文件名（含扩展名），与 eval/questions.jsonl 的 expected_source 对应。
扫描件/图片类 PDF 没有文本层，需 OCR，列为扩展。
"""

import logging
from pathlib import Path

SUPPORTED_EXTS = {".md", ".txt", ".docx", ".xlsx", ".pdf"}

logger = logging.getLogger(__name__)


def _extract_plain_text(path: Path) -> str | None:
    try:
        return path.read_text(encoding="utf-8").strip()
    except (UnicodeDecodeError, OSError):
        return None


def _extract_docx(path: Path) -> str | None:
    """Word：段落 + 表格（表格转 markdown 竖线行）。"""
    try:
        from docx import Document

        doc = Document(str(path))
    except Exception:
        return None
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[表格 {idx}]\n" + "\n".join(rows))
    return "\n\n".join(parts) or None


def _extract_xlsx(path: Path) -> str | None:
    """Excel：每个工作表转成文本行（| 分隔），丢失样式只留内容。

    空格必须**保留占位**而不是丢掉：早先写成 `[str(v) for v in row if v is not None]`，
    于是 `[None, 400, '北京']` 会渲染成 `400 | 北京`——数值 400 落到了"城市"列的位置上，
    表头与数据整体错位。表格里合并/留空单元格是常态，模型据此会把 400 归到错误的列，
    而且答得非常有信心。所以空单元格渲染成空字段，位置一个都不能少。
    """
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        parts = []
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v).strip() for v in row]
                if any(vals):
                    lines.append(" | ".join(vals))
            if lines:
                parts.append(f"[工作表：{ws.title}]\n" + "\n".join(lines))
    finally:
        wb.close()
    return "\n\n".join(parts) or None


def _extract_pdf(path: Path) -> str | None:
    """PDF：逐页抽取文本层，空页跳过；无文本层的扫描件返回内容为空。"""
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
    except Exception:
        return None
    pages = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            pages.append(f"[第 {idx} 页]\n" + text)
    return "\n\n".join(pages) or None


_BINARY_EXTRACTORS = {
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pdf": _extract_pdf,
}


def _extract(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _extract_plain_text(path)
    extractor = _BINARY_EXTRACTORS.get(suffix)
    return extractor(path) if extractor else None


def load_documents(source_dir: str | Path) -> list[dict]:
    """遍历 data/docs 下支持的文档，解析为 [{'content': str, 'source': str}]。

    **只取顶层文件，不递归子目录**（source 用文件名，递归会让同名文件撞车）。
    单个文件解析失败会被跳过（不中断整库建索引），**但每一次跳过都会打日志**——
    原先只在"抽取函数抛异常"时打，而各抽取函数内部把异常吞掉返回 None，
    于是最常见的失败（非 UTF-8 的 .txt、损坏的 docx/xlsx/pdf）是**静默丢文档**：
    语料悄悄变少、检索变差、Agent 答"资料中没有"，却没有任何线索（实测复现）。
    这类静默降级对 RAG 项目是最难查的一类问题，所以宁可吵一点也要每次都说。
    """
    root = Path(source_dir)
    if not root.is_dir():
        raise FileNotFoundError(f"文档目录不存在: {root}")
    documents: list[dict] = []
    skipped: list[str] = []
    for path in sorted(root.iterdir()):
        if path.is_dir():
            if any(
                p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith("_")
                for p in path.rglob("*")
            ):
                logger.warning(
                    "跳过子目录 %s：本加载器只索引 %s 顶层文件，子目录不会被索引",
                    path.name, root,
                )
            continue
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTS:
            continue
        if path.name.startswith("_"):
            continue
        try:
            content = _extract(path)
        except Exception as exc:  # noqa: BLE001
            # 单个文件损坏（例如能打开但中途解析失败的 docx/xlsx/pdf）不应拖垮整库建索引，
            # 否则首次自动建索引会直接变成 503
            logger.warning(
                "跳过无法解析的文件 %s：%s: %s", path.name, type(exc).__name__, exc
            )
            skipped.append(path.name)
            continue
        if content:
            documents.append({"content": content, "source": path.name})
        else:
            # content 为 None：抽取函数内部吞掉了异常（或该格式没有文本层，
            # 例如扫描件 PDF）。空字符串同理——两者都不会进索引，必须显式说出来。
            logger.warning(
                "跳过无文本内容的文件 %s（格式不支持解析、文件损坏，或扫描件 PDF 无文本层）",
                path.name,
            )
            skipped.append(path.name)
    if skipped:
        logger.warning(
            "本次建索引共跳过 %d 个文件：%s", len(skipped), "、".join(skipped)
        )
    if not documents:
        raise ValueError(f"目录下没有可解析的文档: {root}")
    return documents
