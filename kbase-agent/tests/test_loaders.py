"""loader 多格式解析测试：docx/xlsx 运行时生成，pdf 用仓库自带 fixture。"""

from pathlib import Path

from app.retrieval.loader import SUPPORTED_EXTS, load_documents

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _build_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("补卡制度", level=1)
    doc.add_paragraph("忘记打卡须在 3 个工作日内申请补卡。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "限额"
    table.cell(1, 0).text = "上海"
    table.cell(1, 1).text = "500"
    doc.save(str(path))


def _build_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "发薪"
    ws.append(["项目", "说明"])
    ws.append(["发薪日", "每月 10 日"])
    wb.save(str(path))


def test_supported_exts():
    assert {".md", ".txt", ".docx", ".xlsx", ".pdf"} <= SUPPORTED_EXTS


def test_load_docx_xlsx_pdf(tmp_path):
    import shutil

    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    (docs_dir / "_说明.md").write_text("不应被索引", encoding="utf-8")
    (docs_dir / "note.md").write_text("纯文本说明", encoding="utf-8")
    _build_docx(docs_dir / "考勤制度.docx")
    _build_xlsx(docs_dir / "薪酬表.xlsx")
    shutil.copy(FIXTURES / "sample.pdf", docs_dir / "sample.pdf")

    docs = load_documents(docs_dir)
    by_source = {d["source"]: d["content"] for d in docs}

    # _ 前缀文件不参与
    assert "_说明.md" not in by_source

    # docx：段落 + 表格都抽取到
    docx_text = by_source["考勤制度.docx"]
    assert "补卡" in docx_text and "500" in docx_text

    # xlsx：sheet 名 + 单元格内容
    xlsx_text = by_source["薪酬表.xlsx"]
    assert "发薪" in xlsx_text and "每月 10 日" in xlsx_text

    # pdf：文本层逐页抽取
    pdf_text = by_source["sample.pdf"]
    assert "salary" in pdf_text


def test_load_pdf_fixture(tmp_path):
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    # 用仓库 fixture 文件拷贝到临时语料目录
    import shutil

    shutil.copy(FIXTURES / "sample.pdf", docs_dir / "sample.pdf")
    docs = load_documents(docs_dir)
    assert docs and docs[0]["source"] == "sample.pdf"
    assert "10th" in docs[0]["content"]


def test_xlsx_empty_cells_keep_their_column_position(tmp_path):
    """空格必须占位。

    回归背景：原先写成 `[str(v) for v in row if v is not None]`，空格被直接删掉，
    于是 `[None, 400, '北京']` 渲染成 `400 | 北京`——400 落到了"城市"列的位置，
    表头与数据整体错位，模型会把数值归到错误的列上而且答得很自信。
    """
    import openpyxl

    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "差旅"
    ws.append(["城市", "金额", "备注"])
    ws.append([None, 400, "北京"])          # 首列为空
    ws.append(["上海", None, None])          # 后两列为空
    wb.save(str(docs_dir / "差旅表.xlsx"))

    text = {d["source"]: d["content"] for d in load_documents(docs_dir)}["差旅表.xlsx"]
    lines = [ln for ln in text.splitlines() if "|" in ln]
    # 表头 + 两行数据，且列数一致（错位时行内容会塌成两段）
    assert lines[0] == "城市 | 金额 | 备注"
    assert lines[1] == " | 400 | 北京"
    assert lines[2] == "上海 |  | "
    assert all(ln.count("|") == 2 for ln in lines), "每一行的分隔符个数必须一致，否则列就错位了"


def test_loader_logs_skipped_documents(tmp_path, caplog):
    """跳过任何文件都必须留日志——静默丢文档是 RAG 最难查的一类问题。"""
    import logging

    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    (docs_dir / "ok.md").write_text("正常文档", encoding="utf-8")
    # GBK 编码的中文 txt：拷到别的机器上就是这么来的，直读会 UnicodeDecodeError
    (docs_dir / "gbk.txt").write_bytes("中文备注".encode("gbk"))

    with caplog.at_level(logging.WARNING):
        docs = load_documents(docs_dir)

    assert [d["source"] for d in docs] == ["ok.md"]
    assert "gbk.txt" in caplog.text, f"跳过 gbk.txt 没有打日志：{caplog.text!r}"
